import re
import random
import datetime
import pandas as pd
from openpyxl import load_workbook
from risk_scoring import calculate_risk

# =========================
# FILE PATHS
# =========================

OCR_FILE = "OUTPUT/invoice_text.txt"

CSV_FILE = "OUTPUT/extracted_data.csv"

EXCEL_FILE = r"C:\Users\aarsh\OneDrive - IILM University\customs project\extracted_data1.xlsx"


# =========================
# READ OCR TEXT
# =========================

with open(OCR_FILE, "r", encoding="utf-8") as f:
    text = f.read()

print("\n------ OCR TEXT ------\n")
print(text)


# =========================
# EXTRACT FIELDS
# =========================

data = {
    "Invoice Date": "Not Found",
    "Air Waybill": "Not Found",
    "HTS Code": "Not Found",
    "Country": "Not Found",
    "Gross Weight": "Not Found",
    "Net Weight": "Not Found",
    "Invoice Number": "Not Found",
    "Total Value": "Not Found",
    "Shipper": "Not Found",
    "Consignee": "Not Found",
    "Fax": "Not Found"
}


# Invoice Date
date_match = re.search(
    r"(January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{1,2},\s+\d{4}",
    text,
    re.IGNORECASE
)

if date_match:
    data["Invoice Date"] = date_match.group()


# Air Waybill
awb_match = re.search(
    r"(?:AIR WAYBILL|AWB)[^\d]*(\d[\d\s-]{6,})",
    text,
    re.IGNORECASE
)

if awb_match:
    data["Air Waybill"] = awb_match.group(1).strip()


# HTS Code
hts_match = re.search(
    r"(\d{4}\.\d{2})",
    text
)

if hts_match:
    data["HTS Code"] = hts_match.group(1)


# Country
country_match = re.search(
    r"Manufactured in ([A-Za-z\s]+)",
    text,
    re.IGNORECASE
)

if country_match:
    data["Country"] = country_match.group(1).strip()


# Gross Weight
gross_match = re.search(
    r"gross weight[:\s]+([\d.]+)",
    text,
    re.IGNORECASE
)

if gross_match:
    data["Gross Weight"] = gross_match.group(1)


# Net Weight
net_match = re.search(
    r"net weight[:\s]+([\d.]+)",
    text,
    re.IGNORECASE
)

if net_match:
    data["Net Weight"] = net_match.group(1)


# Invoice Number
invoice_match = re.search(
    r"invoice\s*(?:number|no\.?)[:\s]*([A-Za-z0-9\-]+)",
    text,
    re.IGNORECASE
)

if invoice_match:
    data["Invoice Number"] = invoice_match.group(1)


# Total Value
total_match = re.search(
    r"total[:\s]*usd\s*([\d,\.]+)",
    text,
    re.IGNORECASE
)

if total_match:
    data["Total Value"] = total_match.group(1)


# Email -> use as shipper fallback
email_match = re.search(
    r'[\w\.-]+@[\w\.-]+\.\w+',
    text
)

if email_match:
    data["Shipper"] = email_match.group()


# Consignee
consignee_match = re.search(
    r"HFT Electronics",
    text,
    re.IGNORECASE
)

if consignee_match:
    data["Consignee"] = consignee_match.group()

fax_match = re.search(
    r'fax[:\s]*([+\d\s\-\(\)]+)',
    text,
    re.IGNORECASE
)

if fax_match:
    data["Fax"] = fax_match.group(1).strip()


# =========================
# DISPLAY RESULTS
# =========================

print("\nDATA FOUND:")
print(data)

# =========================
# RISK ANALYSIS
# =========================

risk_score, confidence_score, risk_level, decision, risk_reason, country_risk_level = calculate_risk(data)

#shipment status based on decision

if decision == "MANUAL REVIEW":

    shipment_status = "Manual Review"

elif decision == "AUTO APPROVE":

    shipment_status = "Approved"

else:

    shipment_status = "Under Review"

# =========================
# SHIPMENT ID
# =========================

shipment_id = (
    "SHIP-"
    + datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
)
# =========================
# TELEMETRY SIMULATION
# =========================

temperature = random.randint(20, 35)

humidity = random.randint(40, 80)

door_status = random.choice([
    "Closed",
    "Closed",
    "Closed",
    "Open"
])

data["Risk Score"] = risk_score
data["Confidence Score"] = confidence_score
data["Risk Level"] = risk_level
data["Decision"] = decision
data["Risk Reason"] = risk_reason
data["Country Risk Level"] = country_risk_level
data["Shipment Status"] = shipment_status
data["Shipment ID"] = shipment_id
data["Temperature"] = f"{temperature} °C"
data["Humidity"] = f"{humidity} %"
data["Door Status"] = door_status

print("\nRISK ANALYSIS")
print("Risk Score:", risk_score)
print("Confidence Score:", confidence_score)
print("Risk Level:", risk_level)
print("Decision:", decision)
print("Risk Reason:", risk_reason)
print("Country Risk:", country_risk_level)
print("Shipment Status:", shipment_status)
print("Shipment ID:", shipment_id)
print("Temperature:", data["Temperature"])
print("Humidity:", data["Humidity"])
print("Door Status:", data["Door Status"])

# =========================
# SAVE CSV
# =========================

df = pd.DataFrame([data])

df.to_csv(
    CSV_FILE,
    index=False
)


# =========================
# UPDATE EXISTING EXCEL TABLE
# =========================

try:

    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    # Row 2 = current invoice

    ws["A2"] = data["Invoice Date"]
    ws["B2"] = data["Air Waybill"]
    ws["C2"] = data["HTS Code"]
    ws["D2"] = data["Country"]
    ws["E2"] = data["Gross Weight"]
    ws["F2"] = data["Net Weight"]
    ws["G2"] = data["Invoice Number"]
    ws["H2"] = data["Total Value"]
    ws["I2"] = data["Shipper"]
    ws["J2"] = data["Consignee"]
    ws["K2"] = data["Fax"]
    ws["L2"] = data["Risk Score"]
    ws["M2"] = data["Confidence Score"]
    ws["N2"] = data["Risk Level"]
    ws["O2"] = data["Decision"]
    ws["P2"] = data["Risk Reason"]
    ws["Q2"] = data["Country Risk Level"]
    ws["R2"] = data["Shipment Status"]
    ws["S2"] = data["Shipment ID"]
    ws["U2"] = data["Temperature"]
    ws["V2"] = data["Humidity"]
    ws["W2"] = data["Door Status"]
    wb.save(EXCEL_FILE)

    print("\nExcel table updated successfully")

except PermissionError:

    print(
        "\nERROR: Close extracted_data1.xlsx before processing."
    )

except Exception as e:

    print("\nExcel Update Error:")
    print(e)

print("\nPROCESS COMPLETED")